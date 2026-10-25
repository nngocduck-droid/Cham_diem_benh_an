"""Tạo file Excel tổng hợp kết quả chấm (dùng chung cho tải về và gửi email)."""
import io
import re
from thoi_gian import now_vn

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def excel_filename() -> str:
    return f"KetQua_BenhAn_{now_vn().strftime('%d%m%Y_%H%M')}.xlsx"


def _safe_sheet_name(ten_file: str, wb) -> str:
    """Tên sheet Excel: tối đa 31 ký tự, không chứa []:*?/\\ và không trùng."""
    base = re.sub(r"\.pdf$", "", ten_file or "", flags=re.I)
    base = re.sub(r"[\[\]:*?/\\]", "_", base).strip("' ")[:28] or "Benh an"
    name, i = base, 2
    existing = {s.lower() for s in wb.sheetnames}
    while name.lower() in existing:
        suffix = f" ({i})"
        name = base[: 31 - len(suffix)] + suffix
        i += 1
    return name


def build_excel(results: list, plagiarism: list = None) -> bytes:
    wb = openpyxl.Workbook()

    # ── Sheet 1: Tổng hợp ──
    ws1 = wb.active
    ws1.title = "Tổng hợp"

    blue = PatternFill("solid", fgColor="1F6AA5")
    light = PatternFill("solid", fgColor="D6E8F7")
    green_fill = PatternFill("solid", fgColor="E2EFDA")
    red_fill = PatternFill("solid", fgColor="FCE4D6")
    white_bold = Font(name="Times New Roman", bold=True, color="FFFFFF", size=11)
    bold = Font(name="Times New Roman", bold=True, size=11)
    normal = Font(name="Times New Roman", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="AAAAAA")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws1.merge_cells("A1:H1")
    ws1["A1"] = "BẢNG TỔNG HỢP ĐIỂM CHẤM BỆNH ÁN Y KHOA"
    ws1["A1"].font = Font(name="Times New Roman", bold=True, size=14, color="1F6AA5")
    ws1["A1"].alignment = center

    ws1.merge_cells("A2:H2")
    ws1["A2"] = f"Ngày xuất: {now_vn().strftime('%d/%m/%Y %H:%M')}  |  Tổng số bệnh án: {len(results)}"
    ws1["A2"].font = Font(name="Times New Roman", italic=True, size=10, color="666666")
    ws1["A2"].alignment = center
    ws1.append([])

    headers = ["STT", "File bệnh án", "Họ tên BN", "Tuổi", "Giới", "Tổng điểm", "Xếp loại", "Thời gian"]
    ws1.append(headers)
    for col, h in enumerate(headers, 1):
        c = ws1.cell(row=4, column=col)
        c.fill = blue; c.font = white_bold; c.alignment = center; c.border = bd

    for i, r in enumerate(results, 1):
        bn = r.get("benh_nhan", {})
        diem = r.get("tong_diem", 0)
        fill = green_fill if diem >= 5 else red_fill
        row = [i, r.get("ten_file",""), bn.get("ho_ten",""), bn.get("tuoi",""),
               bn.get("gioi",""), diem, r.get("xep_loai",""), r.get("thoi_gian","")]
        ws1.append(row)
        ro = ws1.max_row
        for col in range(1, 9):
            c = ws1.cell(row=ro, column=col)
            c.fill = light if i % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
            c.font = normal; c.border = bd
            c.alignment = center if col != 2 else left
        ws1.cell(ro, 6).fill = fill
        ws1.cell(ro, 6).font = Font(name="Times New Roman", bold=True, size=11)

    ws1.column_dimensions["A"].width = 6
    ws1.column_dimensions["B"].width = 30
    ws1.column_dimensions["C"].width = 20
    ws1.column_dimensions["D"].width = 8
    ws1.column_dimensions["E"].width = 8
    ws1.column_dimensions["F"].width = 12
    ws1.column_dimensions["G"].width = 14
    ws1.column_dimensions["H"].width = 18

    # ── Sheet 2: Chi tiết từng bệnh án ──
    for r in results:
        ws = wb.create_sheet(title=_safe_sheet_name(r.get("ten_file", ""), wb))
        ws.merge_cells("A1:F1")
        ws["A1"] = f"KẾT QUẢ CHẤM: {r.get('ten_file','')}"
        ws["A1"].font = Font(name="Times New Roman", bold=True, size=13, color="1F6AA5")
        ws["A1"].alignment = center

        bn = r.get("benh_nhan", {})
        ws.merge_cells("A2:F2")
        ws["A2"] = f"Bệnh nhân: {bn.get('ho_ten','')} | Tuổi: {bn.get('tuoi','')} | Giới: {bn.get('gioi','')} | Mã: {bn.get('ma_benh_nhan','')}"
        ws["A2"].font = Font(name="Times New Roman", italic=True, size=10)
        ws["A2"].alignment = center
        ws.append([])

        hdrs = ["STT", "Nội dung", "Điểm tối đa", "Điểm đạt", "Tỉ lệ", "Nhận xét"]
        ws.append(hdrs)
        for col, h in enumerate(hdrs, 1):
            c = ws.cell(row=4, column=col)
            c.fill = blue; c.font = white_bold; c.alignment = center; c.border = bd

        for idx, item in enumerate(r.get("chi_tiet", []), 1):
            pct = f"{round(item['diem_dat']/item['diem_toi_da']*100)}%" if item['diem_toi_da'] > 0 else "0%"
            ws.append([item["stt"], item["ten_muc"], item["diem_toi_da"], item["diem_dat"], pct, item["nhan_xet"]])
            ro = ws.max_row
            fill_row = light if idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
            for col in range(1, 7):
                c = ws.cell(row=ro, column=col)
                c.fill = fill_row; c.font = normal; c.border = bd
                c.alignment = center if col in [1,3,4,5] else left

        # Total
        tr = ws.max_row + 1
        ws.merge_cells(f"A{tr}:B{tr}")
        ws.cell(tr,1,"TỔNG ĐIỂM").font = white_bold
        ws.cell(tr,1).fill = blue; ws.cell(tr,1).alignment = center; ws.cell(tr,1).border = bd
        ws.cell(tr,3,10.0).fill = blue; ws.cell(tr,3).font = white_bold; ws.cell(tr,3).alignment = center; ws.cell(tr,3).border = bd
        ws.cell(tr,4, r.get("tong_diem",0)).font = Font(name="Times New Roman", bold=True, size=13, color="C00000")
        ws.cell(tr,4).fill = green_fill; ws.cell(tr,4).alignment = center; ws.cell(tr,4).border = bd
        ws.cell(tr,5, r.get("xep_loai","")).font = bold; ws.cell(tr,5).fill = green_fill; ws.cell(tr,5).alignment = center; ws.cell(tr,5).border = bd
        ws.cell(tr,6, r.get("nhan_xet_chung","")).font = normal; ws.cell(tr,6).fill = green_fill; ws.cell(tr,6).alignment = left; ws.cell(tr,6).border = bd

        # Các phần cần sửa
        phan_sua = r.get("cac_phan_can_sua", [])
        if phan_sua:
            sr = ws.max_row + 1
            ws.merge_cells(f"A{sr}:B{sr}")
            ws.cell(sr, 1, "CÁC PHẦN CẦN SỬA").font = Font(name="Times New Roman", bold=True, size=11, color="C00000")
            ws.cell(sr, 1).fill = PatternFill("solid", fgColor="FCE4D6"); ws.cell(sr, 1).alignment = center; ws.cell(sr, 1).border = bd
            ws.merge_cells(f"C{sr}:F{sr}")
            sua_text = "\n".join(f"{i+1}. {s}" for i, s in enumerate(phan_sua))
            ws.cell(sr, 3, sua_text).font = normal; ws.cell(sr, 3).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True); ws.cell(sr, 3).border = bd
            ws.cell(sr, 3).fill = PatternFill("solid", fgColor="FFF2CC")
            ws.row_dimensions[sr].height = max(15 * len(phan_sua), 30)

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 26
        ws.column_dimensions["C"].width = 13
        ws.column_dimensions["D"].width = 11
        ws.column_dimensions["E"].width = 10
        ws.column_dimensions["F"].width = 55

    # ── Sheet 3: Kiểm tra đạo văn ──
    plagiarism = plagiarism or []
    ws3 = wb.create_sheet(title="Kiểm tra đạo văn")
    ws3.merge_cells("A1:E1")
    ws3["A1"] = "DANH SÁCH CÁC CẶP BỆNH ÁN NGHI VẤN TRÙNG LẶP"
    ws3["A1"].font = Font(name="Times New Roman", bold=True, size=13, color="C00000")
    ws3["A1"].alignment = center
    ws3.append([])

    if plagiarism:
        hdrs3 = ["Sinh viên 1", "Sinh viên 2", "Mức độ", "Độ tương đồng (%)", "Chi tiết"]
        ws3.append(hdrs3)
        for col, h in enumerate(hdrs3, 1):
            c = ws3.cell(row=3, column=col)
            c.fill = PatternFill("solid", fgColor="C00000")
            c.font = white_bold; c.alignment = center; c.border = bd
        for p in plagiarism:
            ws3.append([p["sv1"], p["sv2"], p["muc_do"], p["do_tuong_dong"],
                        " | ".join(p.get("chi_tiet", []))])
            ro = ws3.max_row
            for col in range(1, 6):
                c = ws3.cell(row=ro, column=col)
                c.fill = PatternFill("solid", fgColor="FFF2CC")
                c.font = normal; c.border = bd
                c.alignment = center if col in [3,4] else left
    else:
        ws3.merge_cells("A3:E3")
        ws3["A3"] = "✅ Không phát hiện trùng lặp đáng kể giữa các bệnh án"
        ws3["A3"].font = Font(name="Times New Roman", size=12, color="2D7A4F")
        ws3["A3"].alignment = center

    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 30
    ws3.column_dimensions["C"].width = 14
    ws3.column_dimensions["D"].width = 18
    ws3.column_dimensions["E"].width = 50


    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_admin_excel(records: list, bo_loc: str = "") -> bytes:
    """Excel toàn hệ thống cho admin. records: bản tóm tắt có thêm 'nguoi_cham', 'email_nguoi_cham'."""
    wb = openpyxl.Workbook()
    blue = PatternFill("solid", fgColor="1F6AA5")
    light = PatternFill("solid", fgColor="D6E8F7")
    white = PatternFill("solid", fgColor="FFFFFF")
    green_fill = PatternFill("solid", fgColor="E2EFDA")
    red_fill = PatternFill("solid", fgColor="FCE4D6")
    white_bold = Font(name="Times New Roman", bold=True, color="FFFFFF", size=11)
    normal = Font(name="Times New Roman", size=11)
    bold = Font(name="Times New Roman", bold=True, size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="AAAAAA")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    def title(ws, text, sub, ncols):
        last = openpyxl.utils.get_column_letter(ncols)
        ws.merge_cells(f"A1:{last}1")
        ws["A1"] = text
        ws["A1"].font = Font(name="Times New Roman", bold=True, size=14, color="1F6AA5")
        ws["A1"].alignment = center
        ws.merge_cells(f"A2:{last}2")
        ws["A2"] = sub
        ws["A2"].font = Font(name="Times New Roman", italic=True, size=10, color="666666")
        ws["A2"].alignment = center
        ws.append([])

    def header(ws, cols, row=4):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = blue; cell.font = white_bold; cell.alignment = center; cell.border = bd

    sub = f"Ngày xuất: {now_vn().strftime('%d/%m/%Y %H:%M')}  |  Số bệnh án: {len(records)}"
    if bo_loc:
        sub += f"  |  Bộ lọc: {bo_loc}"

    # ── Sheet 1: tất cả bệnh án ──
    ws = wb.active
    ws.title = "Tổng hợp"
    cols = ["STT", "Người chấm", "Email người chấm", "File bệnh án", "Họ tên BN", "Tuổi", "Giới",
            "Tổng điểm", "Xếp loại", "Thời gian chấm"]
    title(ws, "TỔNG HỢP ĐIỂM CHẤM BỆNH ÁN – TOÀN HỆ THỐNG", sub, len(cols))
    header(ws, cols)
    for i, r in enumerate(records, 1):
        bn = r.get("benh_nhan") or {}
        diem = r.get("tong_diem") or 0
        ws.append([i, r.get("nguoi_cham", ""), r.get("email_nguoi_cham", ""), r.get("ten_file", ""),
                   bn.get("ho_ten", ""), bn.get("tuoi", ""), bn.get("gioi", ""), diem,
                   r.get("xep_loai", ""), r.get("thoi_gian", "")])
        ro = ws.max_row
        for c in range(1, len(cols) + 1):
            cell = ws.cell(ro, c)
            cell.fill = light if i % 2 == 0 else white
            cell.font = normal; cell.border = bd
            cell.alignment = left if c in (2, 3, 4, 5) else center
        ws.cell(ro, 8).fill = green_fill if float(diem or 0) >= 5 else red_fill
        ws.cell(ro, 8).font = bold
    for col, w in zip("ABCDEFGHIJ", [6, 22, 26, 30, 22, 7, 7, 10, 12, 17]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    # ── Sheet 2: theo người chấm ──
    ws2 = wb.create_sheet("Theo người chấm")
    cols2 = ["STT", "Người chấm", "Email", "Số bệnh án", "Điểm TB", "Thấp nhất", "Cao nhất", "Đạt (≥ 5)", "Tỉ lệ đạt"]
    title(ws2, "THỐNG KÊ THEO NGƯỜI CHẤM", sub, len(cols2))
    header(ws2, cols2)
    groups = {}
    for r in records:
        key = (r.get("nguoi_cham", ""), r.get("email_nguoi_cham", ""))
        groups.setdefault(key, []).append(float(r.get("tong_diem") or 0))
    rows = sorted(groups.items(), key=lambda kv: -len(kv[1]))
    for i, ((ten, email), diems) in enumerate(rows, 1):
        dat = sum(1 for d in diems if d >= 5)
        ws2.append([i, ten, email, len(diems), round(sum(diems) / len(diems), 2), min(diems), max(diems),
                    dat, f"{round(dat / len(diems) * 100)}%"])
        ro = ws2.max_row
        for c in range(1, len(cols2) + 1):
            cell = ws2.cell(ro, c)
            cell.fill = light if i % 2 == 0 else white
            cell.font = normal; cell.border = bd
            cell.alignment = left if c in (2, 3) else center
    for col, w in zip("ABCDEFGHI", [6, 24, 28, 12, 10, 11, 11, 11, 11]):
        ws2.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def admin_excel_filename() -> str:
    return f"TongHop_ToanHeThong_{now_vn().strftime('%d%m%Y_%H%M')}.xlsx"
